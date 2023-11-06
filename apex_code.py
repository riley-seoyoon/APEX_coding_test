#!/usr/bin/env python3

import os
import math
import statistics
import numpy as np
import scipy.stats
import pandas as pd
import docx

os.chdir("D:\APEX_coding_test")

# Read excel file into dataframe
data = pd.read_excel("APEX_input.xlsx").set_index("group_num")

group = data.index.unique()

# Separate day 0 and 13
day0 = data[data["days_after_treatment"] == 0]
day13 = data[data["days_after_treatment"] == 13]

# Calculate means and standard deviations for day 0 and day 13 for each group_num
day0_summary = day0.groupby('group_num').agg(['mean', 'std'])
day13_summary = day13.groupby('group_num').agg(['mean', 'std'])

# Calculate TGI
control_change = day13_summary.loc[1, ('volume_mm3', 'mean')] - day0_summary.loc[1, ('volume_mm3', 'mean')]
tgi = (1 - ((day13_summary['volume_mm3']['mean'] - day0_summary['volume_mm3']['mean']) / control_change)) * 100

# Join values into one dataframe
joined_summary_0 = day0_summary.apply(lambda x: f"{x.volume_mm3['mean']:.2f} ± {x.volume_mm3['std']:.2f}", axis=1)
joined_summary_13 = day13_summary.apply(lambda x: f"{x.volume_mm3['mean']:.2f} ± {x.volume_mm3['std']:.2f}", axis=1)
df = pd.concat([joined_summary_0, joined_summary_13, tgi], axis = 1).reset_index()
df = df.rename(columns = {"group_num": "Group", 0 : "Day 0 \nmean ± SD", 1: "Day 13 \nmean ± SD", "mean": "TGI(%)"})

# Create word document and table
document = docx.Document()
table = document.add_table(rows=1, cols=len(df.columns))

# Add column names to the first row
for i, col_name in enumerate(df.columns):
    table.cell(0, i).text = col_name

# Add values into table
values = df.values
for i in range(df.shape[0]):
    row = table.add_row()
    for j, cell in enumerate(row.cells):
        cell.text = str(values[i, j]) 

table.cell(1, 0).text = '1 (control)'
table.style = "Table Grid"

# Save document
document.save("result.docx")